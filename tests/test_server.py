"""Tests for the FastAPI server layer (server/) - the Gradio replacement.

Covers input validation (400/413), the SSE streaming contract of both pipeline
steps against the REAL graphs wired with fake LLMs (no network, no API cost),
pipeline queue/serialization semantics (Gradio queue port), artifact download
guards, app-wide basic auth, and health reporting.

Run with: venv\\Scripts\\python.exe -m pytest tests/test_server.py -v
"""

import json
import os
import sys
import threading
import uuid
from pathlib import Path

import pytest

os.environ["LANGCHAIN_TRACING_V2"] = "false"
os.environ.setdefault("GROQ_API_KEY", "gsk_test_dummy_key_for_tests")

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402
from langchain_core.messages import AIMessage  # noqa: E402

import agent  # noqa: E402
import agent.config  # noqa: E402
import agent.llms  # noqa: E402
import agent.tracing  # noqa: E402
import server.app as server_app  # noqa: E402
import server.runners as runners  # noqa: E402

JD = "Python engineer with Docker and Kubernetes experience."
RESUME_TEXT = "Jane Roe\nPython developer with cloud exposure."


# ==========================================
# Shared fakes (same shape as test_resume_agent)
# ==========================================

class FakeStrictLLM:
    def invoke(self, messages):
        return AIMessage(content="- Kubernetes\n- Terraform")

    def with_structured_output(self, schema):
        class Structured:
            def invoke(self, messages):
                return agent.ReviewOutput(score=85, feedback="Add metrics to bullets.")
        return Structured()


class FakeCreativeLLM:
    def invoke(self, messages):
        return AIMessage(content="## Skills\nPython, Docker\n## Experience\nBuilt ML pipelines.")


class FakeFastLLM:
    def invoke(self, messages):
        prompt = messages[-1].content
        if "Cover Letter" in prompt:
            content = "Dear Hiring Team,\nI am excited to apply.\nSincerely, Candidate"
        else:
            content = "1. Sample question?\n\n## Preparation Tips\n- Tip one\n- Tip two"
        return AIMessage(content=content)

    def with_structured_output(self, schema):
        class Structured:
            def invoke(self, messages):
                return agent.SemanticReview(verdicts=[])
        return Structured()


@pytest.fixture
def fake_llms(monkeypatch):
    monkeypatch.setattr(agent.llms, "llm_strict", FakeStrictLLM())
    monkeypatch.setattr(agent.llms, "llm_creative", FakeCreativeLLM())
    monkeypatch.setattr(agent.llms, "llm_fast", FakeFastLLM())


@pytest.fixture
def clean_name_register(monkeypatch):
    monkeypatch.setattr(agent.tracing, "_KNOWN_NAMES", [])
    monkeypatch.setattr(agent.tracing, "_NAME_MASKS", None)


@pytest.fixture
def client():
    # Without a context manager, lifespan does NOT run: app.state.health stays
    # untouched and the cleanup daemon and startup health check are skipped.
    return TestClient(server_app.app)


def parse_sse(text: str) -> list[tuple[str, dict]]:
    """Parses the text/event-stream protocol into (event, payload) pairs."""
    events = []
    for block in text.split("\n\n"):
        block = block.strip("\r\n")
        if not block:
            continue
        event, data = None, None
        for line in block.splitlines():
            if line.startswith("event: "):
                event = line[len("event: "):].strip()
            elif line.startswith("data: "):
                data = json.loads(line[len("data: "):])
        if event and data is not None:
            events.append((event, data))
    return events


def _statuses(events) -> list[str]:
    return [e.get("message", "") for kind, e in events if kind == "status"]


# ==========================================
# Health
# ==========================================

def test_health_reports_skipped_without_startup_check(client):
    body = client.get("/api/v1/health").json()
    assert body["ok"] is None
    assert "skipped" in body["message"]


def test_health_reflects_startup_check_when_lifespan_ran(monkeypatch):
    monkeypatch.setattr(agent.llms, "check_provider_health",
                        lambda: (False, "groq is down"))
    with TestClient(server_app.app) as c:
        body = c.get("/api/v1/health").json()
    assert body == {"ok": False, "message": "groq is down"}


# ==========================================
# Runners: config + resolution primitives
# ==========================================

def test_thread_config_pins_thread_and_run_id():
    cfg = runners._thread_config("t1", run_id="rid-123")
    assert cfg["configurable"]["thread_id"] == "t1"
    assert cfg["run_id"] == "rid-123"
    assert "callbacks" in cfg

    cfg = runners._thread_config("t2")
    assert "run_id" not in cfg


def test_resolve_resume_source_file_wins_over_text():
    assert runners.resolve_resume_source("p.pdf", "text") == "p.pdf"
    assert runners.resolve_resume_source(None, "  text  ") == "text"
    assert runners.resolve_resume_source(None, "   ") is None


# ==========================================
# Validation (HTTP level, before streaming)
# ==========================================

def test_step1_requires_job_description(client):
    r = client.post("/api/v1/step1",
                    data={"job_description": "   ", "resume_text": RESUME_TEXT})
    assert r.status_code == 400
    assert "Job Description" in r.json()["detail"]


def test_step1_requires_resume(client):
    r = client.post("/api/v1/step1",
                    data={"job_description": JD, "resume_text": "  "})
    assert r.status_code == 400
    assert "resume" in r.json()["detail"].lower()


def test_step1_rejects_unsupported_file_type(client):
    r = client.post("/api/v1/step1",
                    data={"job_description": JD, "resume_text": ""},
                    files={"resume_file": ("resume.txt", b"plain text", "text/plain")})
    assert r.status_code == 400
    assert "PDF or DOCX" in r.json()["detail"]


def test_step1_rejects_oversized_upload(client, monkeypatch):
    monkeypatch.setattr(agent.config, "MAX_PDF_BYTES", 64)
    r = client.post("/api/v1/step1",
                    data={"job_description": JD, "resume_text": ""},
                    files={"resume_file": ("resume.pdf", b"A" * 100, "application/pdf")})
    assert r.status_code == 413


def test_step2_rejects_invalid_session(client):
    r = client.post("/api/v1/step2",
                    data={"job_description": JD, "resume_text": RESUME_TEXT,
                          "session": "not-a-uuid"})
    assert r.status_code == 400
    assert "session" in r.json()["detail"].lower()


# ==========================================
# Step 1: SSE stream over the real graph
# ==========================================

def test_step1_stream_premasks_name_and_returns_gaps(client, fake_llms, clean_name_register):
    r = client.post("/api/v1/step1",
                    data={"job_description": JD, "resume_text": RESUME_TEXT})
    assert r.status_code == 200
    events = parse_sse(r.text)
    kinds = [kind for kind, _ in events]
    assert kinds[-1] == "done"
    assert "error" not in kinds

    done = events[-1][1]
    uuid.UUID(done["session"])  # well-formed uuid
    assert "Kubernetes" in done["missing_skills"]
    assert any("(step 1 of 2)" in s for s in _statuses(events))
    assert any("Loading documents..." in s for s in _statuses(events))
    # node-aware statuses: every tagged event carries a node id + running state
    tagged = [e for kind, e in events if kind == "status" and e.get("node")]
    assert tagged and all(e["node_state"] == "running" for e in tagged)
    tagged_nodes = {e["node"] for e in tagged}
    assert {"loader", "scanner"} <= tagged_nodes or tagged_nodes == {"loader"}
    assert "Jane Roe" in agent.tracing._KNOWN_NAMES  # pasted resume pre-masked


def test_step1_accepts_docx_upload(client, fake_llms, clean_name_register):
    import io
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Jane Roe")
    doc.add_paragraph("Python developer")
    doc.save(buf)
    r = client.post("/api/v1/step1",
                    data={"job_description": JD, "resume_text": ""},
                    files={"resume_file": ("resume.docx", buf.getvalue(),
                                           "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200
    events = parse_sse(r.text)
    assert events[-1][0] == "done"
    # the loader registers the candidate name parsed from the document
    assert "Jane Roe" in agent.tracing._KNOWN_NAMES


# ==========================================
# Step 2: SSE stream - fresh / resume / reuse flows
# ==========================================

def _spy_feedback(monkeypatch) -> list:
    calls = []
    monkeypatch.setattr(runners, "log_pipeline_feedback",
                        lambda state, run_id: calls.append((state, run_id)))
    return calls


def test_step2_from_scratch_warns_and_returns_full_result(client, fake_llms, monkeypatch):
    feedback_calls = _spy_feedback(monkeypatch)
    r = client.post("/api/v1/step2",
                    data={"job_description": JD, "resume_text": RESUME_TEXT, "notes": ""})
    assert r.status_code == 200
    events = parse_sse(r.text)
    kinds = [kind for kind, _ in events]
    assert kinds[-1] == "result"
    assert "error" not in kinds

    warnings = [e for kind, e in events if kind == "warning"]
    assert warnings and "auto-detected" in warnings[0]["message"]

    result = events[-1][1]
    assert result["optimized_resume"].strip()
    assert result["cover_letter"].startswith("Dear")
    assert "Preparation Tips" in result["interview_questions"]
    assert result["ats_result"].strip()

    expected = round(agent.config.ATS_SCORE_WEIGHT * result["ats_percentage"]
                     + agent.config.LLM_SCORE_WEIGHT * 85)
    assert result["score"] == expected
    assert result["score_class"] == ("good" if expected >= agent.config.SCORE_THRESHOLD
                                     else "warn" if expected >= 60 else "bad")
    assert result["weights"] == {"ats": agent.config.ATS_SCORE_WEIGHT,
                                 "llm": agent.config.LLM_SCORE_WEIGHT}
    assert "composite score" in result["status"]

    artifact_kinds = {a["kind"] for a in result["artifacts"]}
    assert {"resume-pdf", "cover-pdf"} <= artifact_kinds
    for artifact in result["artifacts"]:
        dl = client.get("/" + artifact["url"])
        assert dl.status_code == 200
        assert len(dl.content) > 0

    assert len(feedback_calls) == 1
    state_sent, run_id_sent = feedback_calls[0]
    uuid.UUID(run_id_sent)
    assert state_sent.get("score") == expected

    statuses = _statuses(events)
    assert any("(1/2)" in s for s in statuses)
    assert any("(2/2)" in s for s in statuses)

    # node ids accompany the per-node statuses across the whole run incl. fan-out
    tagged = [e for kind, e in events if kind == "status" and e.get("node")]
    assert tagged and all(e["node_state"] == "running" for e in tagged)
    tagged_nodes = {e["node"] for e in tagged}
    assert {"loader", "scanner", "improver", "reviewer", "ats_check",
            "cover_letter", "interview_prep", "pdf_exporter"} <= tagged_nodes


def test_step2_resumes_interrupted_session_from_step1(client, fake_llms, clean_name_register):
    r1 = client.post("/api/v1/step1",
                     data={"job_description": JD, "resume_text": RESUME_TEXT})
    session = parse_sse(r1.text)[-1][1]["session"]

    r2 = client.post("/api/v1/step2",
                     data={"job_description": JD, "resume_text": RESUME_TEXT,
                           "notes": "I used k8s at work", "session": session})
    events = parse_sse(r2.text)
    assert events[-1][0] == "result"
    result = events[-1][1]
    assert result["session"] == session
    statuses = _statuses(events)
    # checkpoint interruption: segment 1 (loader/scanner) is NOT re-run
    assert not any("(1/2)" in s for s in statuses)
    assert any("(2/2)" in s for s in statuses)
    tagged_nodes = {e["node"] for kind, e in events
                    if kind == "status" and e.get("node")}
    assert "improver" in tagged_nodes
    assert not {"loader", "scanner"} & tagged_nodes

    # the user's edited notes were injected at the human-review checkpoint
    snapshot = agent.agent_app.get_state({"configurable": {"thread_id": session}})
    assert snapshot.values.get("human_notes") == "I used k8s at work"


def test_step2_reuses_completed_session_by_starting_fresh(client, fake_llms, monkeypatch):
    _spy_feedback(monkeypatch)
    r1 = client.post("/api/v1/step2",
                     data={"job_description": JD, "resume_text": RESUME_TEXT,
                           "notes": "some notes"})
    first_session = parse_sse(r1.text)[-1][1]["session"]

    r2 = client.post("/api/v1/step2",
                     data={"job_description": JD, "resume_text": RESUME_TEXT,
                           "notes": "some notes", "session": first_session})
    events = parse_sse(r2.text)
    fresh_events = [e for kind, e in events if kind == "session"]
    assert fresh_events and fresh_events[0]["session"] != first_session
    assert any("fresh run" in s for s in _statuses(events))
    assert events[-1][1]["session"] != first_session


def test_step2_pipeline_error_becomes_error_event(client, fake_llms, monkeypatch, clean_name_register):
    class BrokenImproverLLM:
        def invoke(self, messages):
            raise RuntimeError("provider outage")

    # the improver is the first node after the interrupt -> fails the resume
    monkeypatch.setattr(agent.llms, "llm_creative", BrokenImproverLLM())
    r = client.post("/api/v1/step2",
                    data={"job_description": JD, "resume_text": RESUME_TEXT,
                          "notes": "notes"})
    events = parse_sse(r.text)
    assert events[-1][0] == "error"
    assert "provider outage" in events[-1][1]["message"]


# ==========================================
# Queue semantics (Gradio queue port)
# ==========================================

def test_queue_full_emits_error_event_not_http(client, fake_llms, monkeypatch):
    monkeypatch.setattr(server_app, "_QUEUE_MAX", 0)
    r = client.post("/api/v1/step1",
                    data={"job_description": JD, "resume_text": RESUME_TEXT})
    events = parse_sse(r.text)
    assert len(events) == 1
    kind, payload = events[0]
    assert kind == "error"
    assert "queue is full" in payload["message"]


def test_run_waits_for_free_slot_when_pipeline_busy(client, fake_llms):
    assert server_app._RUN_LOCK.acquire(blocking=False)
    timer = threading.Timer(0.6, server_app._RUN_LOCK.release)
    timer.start()
    try:
        r = client.post("/api/v1/step1",
                        data={"job_description": JD, "resume_text": RESUME_TEXT})
    finally:
        timer.cancel()
    events = parse_sse(r.text)
    assert any(s.startswith("Waiting for a free pipeline slot")
               for s in _statuses(events))
    assert events[-1][0] == "done"  # completed after the slot freed


# ==========================================
# Artifact downloads
# ==========================================

def test_artifact_download_roundtrip(client, monkeypatch, tmp_path):
    monkeypatch.setattr(agent.config, "OUTPUT_ROOT", tmp_path)
    session = str(uuid.uuid4())
    session_dir = tmp_path / session
    session_dir.mkdir()
    (session_dir / "resume_optimized.pdf").write_bytes(b"%PDF-1.4 resume")
    (session_dir / "cover_letter.pdf").write_bytes(b"%PDF-1.4 cover")
    (session_dir / "resume.docx").write_bytes(b"PK\x03\x04 docx bytes")
    (session_dir / "resume.tex").write_text("\\documentclass{article}")

    pdf = client.get(f"/api/v1/artifacts/{session}/resume-pdf",
                     params={"file": "resume_optimized.pdf"})
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert "resume_optimized.pdf" in pdf.headers["content-disposition"]
    assert pdf.content.startswith(b"%PDF")

    cover = client.get(f"/api/v1/artifacts/{session}/cover-pdf",
                       params={"file": "cover_letter.pdf"})
    assert cover.status_code == 200

    docx = client.get(f"/api/v1/artifacts/{session}/resume-docx",
                      params={"file": "resume.docx"})
    assert docx.status_code == 200
    assert docx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    tex = client.get(f"/api/v1/artifacts/{session}/resume-tex",
                     params={"file": "resume.tex"})
    assert tex.status_code == 200
    assert tex.headers["content-type"].startswith("text/x-tex")


def test_artifact_unknown_session_or_kind_is_404(client, monkeypatch, tmp_path):
    monkeypatch.setattr(agent.config, "OUTPUT_ROOT", tmp_path)
    session = str(uuid.uuid4())
    (tmp_path / session).mkdir()

    assert client.get(f"/api/v1/artifacts/{session}/resume-pdf",
                      params={"file": "missing.pdf"}).status_code == 404
    assert client.get(f"/api/v1/artifacts/not-a-uuid/resume-pdf",
                      params={"file": "x.pdf"}).status_code == 404
    assert client.get(f"/api/v1/artifacts/{session}/resume-xls",
                      params={"file": "resume.pdf"}).status_code == 404


def test_artifact_traversal_is_contained(client, monkeypatch, tmp_path):
    monkeypatch.setattr(agent.config, "OUTPUT_ROOT", tmp_path)
    session = str(uuid.uuid4())
    (tmp_path / session).mkdir()
    (tmp_path / "outside_secret.pdf").write_bytes(b"%PDF")

    # Path.name strips the '../' component, so the probe looks for a
    # non-existent file inside the session dir and is contained by construction.
    r = client.get(f"/api/v1/artifacts/{session}/resume-pdf",
                   params={"file": "../outside_secret.pdf"})
    assert r.status_code == 404


# ==========================================
# App-wide basic auth
# ==========================================

def test_auth_required_when_configured(client, monkeypatch):
    monkeypatch.setattr(server_app, "AUTH_CREDENTIALS", ("admin", "s3cret"))
    r = client.get("/api/v1/health")
    assert r.status_code == 401
    assert "basic" in r.headers["www-authenticate"].lower()

    r_wrong = client.get("/api/v1/health", auth=("admin", "nope"))
    assert r_wrong.status_code == 401

    r_ok = client.get("/api/v1/health", auth=("admin", "s3cret"))
    assert r_ok.status_code == 200


def test_api_documentation_not_exposed(client, monkeypatch):
    # FastAPI's built-in docs don't inherit app-level auth dependencies, so
    # the schema routes are disabled entirely (internal API; the BFF owns
    # the public contract).
    assert client.get("/openapi.json").status_code == 404
    assert client.get("/docs").status_code == 404


def test_open_access_when_auth_not_configured(client):
    assert client.get("/api/v1/health").status_code == 200
